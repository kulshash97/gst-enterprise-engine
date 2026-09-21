import io
from datetime import datetime
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from core.models import Invoice, InvoiceItem

def fetch_period_invoices(db, tenant_id, period):
    prefix = f"{period[2:]}-{period[:2]}"
    return db.query(Invoice).filter(Invoice.tenant_id == tenant_id, Invoice.doc_date.like(f"{prefix}%")).all()

def generate_gstr1_json_payload(db, tenant, period):
    invoices = fetch_period_invoices(db, tenant.id, period)
    b2b_grouped = defaultdict(list)
    hsn_summary = defaultdict(lambda: {"desc": "", "uqc": "NOS", "qty": 0.0, "val": 0.0, "txval": 0.0, "iamt": 0.0, "camt": 0.0, "samt": 0.0})

    for inv in invoices:
        items = db.query(InvoiceItem).filter(InvoiceItem.invoice_id == inv.id).all()
        is_inter = (inv.supplier_state_code != inv.buyer_state_code)
        dt_str = datetime.strptime(inv.doc_date, "%Y-%m-%d").strftime("%d-%m-%Y")
        
        item_details = []
        for idx, itm in enumerate(items, 1):
            rate = itm.tax_rate
            txval = itm.taxable_value
            igst = round(txval * (rate / 100.0), 2) if is_inter else 0.0
            cgst = round((txval * (rate / 100.0)) / 2.0, 2) if not is_inter else 0.0
            sgst = round((txval * (rate / 100.0)) / 2.0, 2) if not is_inter else 0.0

            item_details.append({"num": idx, "itm_det": {"rt": rate, "txval": txval, "iamt": igst, "camt": cgst, "samt": sgst, "csamt": 0.0}})

            h = hsn_summary[itm.hsn_code]
            h["desc"] = itm.item_name
            h["qty"] += itm.quantity
            h["txval"] += txval
            h["val"] += (txval + igst + cgst + sgst)
            h["iamt"] += igst
            h["camt"] += cgst
            h["samt"] += sgst

        b2b_grouped[inv.buyer_gstin].append({
            "inum": inv.doc_number, "idt": dt_str, "val": inv.grand_total, "pos": inv.buyer_state_code,
            "rchrg": "N", "inv_typ": "R", "itms": item_details
        })

    b2b_list = [{"ctin": ctin, "inv": invs} for ctin, invs in b2b_grouped.items()]
    hsn_list = [
        {"num": i, "hsn_sc": k, "desc": v["desc"], "uqc": "NOS", "qty": round(v["qty"], 2), "val": round(v["val"], 2),
         "txval": round(v["txval"], 2), "iamt": round(v["iamt"], 2), "camt": round(v["camt"], 2), "samt": round(v["samt"], 2), "csamt": 0.0}
        for i, (k, v) in enumerate(hsn_summary.items(), 1)
    ]

    return {
        "gstin": invoices[0].supplier_gstin if invoices else "36AAACG1234F1Z5",
        "fp": period, "gt": 0.0, "cur_gt": 0.0, "b2b": b2b_list, "hsn": {"data": hsn_list}
    }

def generate_gstr1_excel_workbook(db, tenant, period):
    invoices = fetch_period_invoices(db, tenant.id, period)
    wb = Workbook()
    h_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    h_font = Font(name="Arial", size=9, bold=True, color="FFFFFF")

    ws_b2b = wb.active
    ws_b2b.title = "B2B"
    b2b_cols = ["Recipient GSTIN", "Receiver Name", "Invoice No", "Date", "Total Value", "POS", "Reverse Charge", "Rate", "Taxable Value"]
    ws_b2b.append(b2b_cols)
    for c in ws_b2b[1]:
        c.fill = h_fill
        c.font = h_font

    for inv in invoices:
        for itm in db.query(InvoiceItem).filter(InvoiceItem.invoice_id == inv.id).all():
            ws_b2b.append([inv.buyer_gstin, inv.buyer_legal_name, inv.doc_number, inv.doc_date, inv.grand_total, inv.buyer_state_code, "N", itm.tax_rate, itm.taxable_value])

    ws_hsn = wb.create_sheet(title="HSN")
    ws_hsn.append(["HSN", "Description", "UQC", "Total Qty", "Total Value", "Taxable Value"])
    for c in ws_hsn[1]:
        c.fill = h_fill
        c.font = h_font

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
