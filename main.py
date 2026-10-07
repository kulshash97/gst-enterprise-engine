import sys
import os
import secrets
import traceback
import hashlib
import io
from pathlib import Path
from datetime import datetime, date
from typing import Optional, List, Dict, Any

sys.path.insert(0, str(Path(__file__).resolve().parent))

from fastapi import FastAPI, Depends, HTTPException, Header, Query, BackgroundTasks, Request, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, StreamingResponse, JSONResponse, Response
from pydantic import BaseModel
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from core.database import Base, engine, get_db
from core.models import Tenant, Invoice, InvoiceItem, ComplianceAuditLog
from schemas.invoice_schema import CreateInvoiceRequest
from compliance.tax_engine import compute_gst_split
from compliance.nic_schema import build_nic_standard_payload
from services.pdf_generator import build_gst_tax_invoice_pdf
from services.email_service import send_invoice_email_task
from services.gstr1_exporter import generate_gstr1_json_payload, generate_gstr1_excel_workbook
from services.reconciliation_engine import (
    run_gstr2b_reconciliation,
    parse_purchase_register_file,
    parse_gstr2b_json_file
)

Base.metadata.create_all(bind=engine)

# Auto-seed authorized tenant
db_init = next(get_db())
try:
    existing_tenant = db_init.query(Tenant).filter(Tenant.api_key == "gst_EgKLOK0WOPPoQHjJ7hPJGU334kjDsvnTpVY53iyM8jc").first()
    if not existing_tenant:
        db_init.add(Tenant(org_name="apex", api_key="gst_EgKLOK0WOPPoQHjJ7hPJGU334kjDsvnTpVY53iyM8jc"))
        db_init.commit()
finally:
    db_init.close()

app = FastAPI(title="ApexTax Enterprise GST Engine", version="2.5.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

os.makedirs("static", exist_ok=True)
app.mount("/static", StaticFiles(directory="static"), name="static")

@app.exception_handler(Exception)
async def global_exception_handler(request: Request, exc: Exception):
    traceback.print_exc()
    return JSONResponse(status_code=500, content={"detail": f"{type(exc).__name__}: {str(exc)}"})

def verify_tenant(
    x_api_key_header: Optional[str] = Header(None, alias="x-api-key"),
    x_api_key_query: Optional[str] = Query(None, alias="x-api-key"),
    db: Session = Depends(get_db)
) -> Tenant:
    key = x_api_key_header or x_api_key_query
    if not key:
        raise HTTPException(status_code=401, detail="API Key missing. Provide via 'x-api-key' header or query parameter.")
    tenant = db.query(Tenant).filter(Tenant.api_key == key.strip()).first()
    if not tenant:
        raise HTTPException(status_code=401, detail="Invalid API Key. Unauthorized access.")
    return tenant

@app.get("/")
def serve_dashboard():
    return FileResponse("static/index.html")

@app.post("/api/v1/tenants/register")
def register_tenant(org_name: str, db: Session = Depends(get_db)):
    api_key = f"gst_{secrets.token_urlsafe(32)}"
    new_tenant = Tenant(org_name=org_name, api_key=api_key)
    db.add(new_tenant)
    db.commit()
    db.refresh(new_tenant)
    return {"status": "TENANT_REGISTERED", "tenant_id": new_tenant.id, "organization": new_tenant.org_name, "api_key": new_tenant.api_key}

@app.post("/api/v1/invoices/process")
def process_invoice(
    payload: CreateInvoiceRequest,
    background_tasks: BackgroundTasks,
    tenant: Tenant = Depends(verify_tenant),
    db: Session = Depends(get_db)
):
    try:
        if payload.supplier.gstin == payload.recipient.gstin:
            raise HTTPException(status_code=400, detail="Supplier and Recipient GSTIN cannot be identical.")

        tax_data = compute_gst_split(payload.supplier.gstin, payload.recipient.gstin, payload.items)
        grand_total = float(tax_data["grand_total"])
        eway_required = grand_total >= 50000.0
        eway_status = "NOT_REQUIRED"
        if eway_required:
            eway_status = "READY_FOR_DISPATCH" if payload.vehicle_number else "PENDING_PART_B_VEHICLE"

        irn, nic_payload, qr_data = build_nic_standard_payload(
            supplier=payload.supplier,
            recipient=payload.recipient,
            tax_data=tax_data,
            doc_number=payload.doc_number,
            doc_date=payload.doc_date
        )

        existing = db.query(Invoice).filter(
            Invoice.tenant_id == tenant.id,
            Invoice.doc_number == payload.doc_number
        ).first()

        if existing:
            db.query(InvoiceItem).filter(InvoiceItem.invoice_id == existing.id).delete()
            db.delete(existing)
            db.commit()

        db_inv = Invoice(
            tenant_id=tenant.id,
            doc_number=payload.doc_number,
            doc_date=str(payload.doc_date),
            supplier_legal_name=payload.supplier.legal_name,
            supplier_gstin=payload.supplier.gstin,
            supplier_state_code=payload.supplier.state_code,
            buyer_legal_name=payload.recipient.legal_name,
            buyer_gstin=payload.recipient.gstin,
            buyer_state_code=payload.recipient.state_code,
            total_taxable_value=float(tax_data["total_taxable_value"]),
            cgst_amount=float(tax_data["total_cgst"]),
            sgst_amount=float(tax_data["total_sgst"]),
            igst_amount=float(tax_data["total_igst"]),
            grand_total=grand_total,
            irn=irn,
            eway_bill_required=eway_required,
            eway_bill_status=eway_status,
            vehicle_number=payload.vehicle_number or ""
        )
        db.add(db_inv)
        db.flush()

        for itm in tax_data["line_items"]:
            db.add(InvoiceItem(
                invoice_id=db_inv.id,
                item_name=itm["item_name"],
                hsn_code=itm["hsn_code"],
                quantity=float(itm["quantity"]),
                unit_price=float(itm["unit_price"]),
                taxable_value=float(itm["taxable_value"]),
                tax_rate=float(itm["tax_rate"])
            ))

        db.add(ComplianceAuditLog(
            tenant_id=tenant.id,
            event_type="INVOICE_GENERATED_AND_DISPATCHED",
            doc_number=payload.doc_number,
            details=f"Total: Rs. {grand_total:,.2f} | IRN: {irn[:16]}..."
        ))
        db.commit()

        if payload.recipient.email:
            try:
                background_tasks.add_task(
                    send_invoice_email_task,
                    recipient_email=payload.recipient.email,
                    doc_number=payload.doc_number,
                    grand_total=grand_total
                )
            except Exception:
                pass

        return {
            "status": "SUCCESS",
            "doc_number": payload.doc_number,
            "email_status": f"QUEUED_FOR_{payload.recipient.email}" if payload.recipient.email else "NO_EMAIL",
            "tax_summary": {
                "is_interstate": bool(tax_data["is_interstate"]),
                "total_taxable_value": float(tax_data["total_taxable_value"]),
                "cgst": float(tax_data["total_cgst"]),
                "sgst": float(tax_data["total_sgst"]),
                "igst": float(tax_data["total_igst"]),
                "grand_total": grand_total
            },
            "compliance": {
                "irn": irn,
                "eway_bill_status": eway_status,
                "vehicle_number": payload.vehicle_number or "N/A"
            }
        }
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/api/v1/invoices/{doc_number:path}/pdf")
def download_invoice_pdf(doc_number: str, tenant: Tenant = Depends(verify_tenant), db: Session = Depends(get_db)):
    inv = db.query(Invoice).filter(Invoice.tenant_id == tenant.id, Invoice.doc_number == doc_number).first()
    if not inv:
        raise HTTPException(status_code=404, detail="Invoice not found.")

    items = db.query(InvoiceItem).filter(InvoiceItem.invoice_id == inv.id).all()
    is_inter = (inv.supplier_state_code != inv.buyer_state_code)

    pdf_data = {
        "doc_number": inv.doc_number,
        "doc_date": inv.doc_date,
        "supplier": {
            "legal_name": inv.supplier_legal_name,
            "gstin": inv.supplier_gstin,
            "address": "Plot 42, Hitec City",
            "city": "Hyderabad",
            "state_code": inv.supplier_state_code,
            "pincode": 500081
        },
        "recipient": {
            "legal_name": inv.buyer_legal_name,
            "gstin": inv.buyer_gstin,
            "address": "Brigade Gateway",
            "city": "Bengaluru",
            "state_code": inv.buyer_state_code,
            "pincode": 560055
        },
        "tax_summary": {
            "is_interstate": is_inter,
            "total_taxable_value": inv.total_taxable_value,
            "cgst": inv.cgst_amount,
            "sgst": inv.sgst_amount,
            "igst": inv.igst_amount,
            "grand_total": inv.grand_total
        },
        "compliance": {
            "e_invoice": {
                "irn": inv.irn,
                "signed_qr_data": {
                    "IRN": inv.irn,
                    "Seller": inv.supplier_gstin,
                    "DocNo": inv.doc_number,
                    "Total": inv.grand_total
                }
            },
            "e_way_bill": {
                "status": inv.eway_bill_status,
                "vehicle_number": inv.vehicle_number or "N/A"
            }
        },
        "line_items": [
            {
                "item_name": it.item_name,
                "hsn_code": it.hsn_code,
                "quantity": it.quantity,
                "unit_price": it.unit_price,
                "taxable_value": it.taxable_value,
                "tax_rate": it.tax_rate,
                "cgst": round((it.taxable_value * (it.tax_rate / 100)) / 2, 2) if not is_inter else 0.0,
                "sgst": round((it.taxable_value * (it.tax_rate / 100)) / 2, 2) if not is_inter else 0.0,
                "igst": round(it.taxable_value * (it.tax_rate / 100), 2) if is_inter else 0.0,
                "line_total": round(it.taxable_value * (1 + it.tax_rate / 100), 2)
            }
            for it in items
        ]
    }
    pdf_buf = build_gst_tax_invoice_pdf(pdf_data)
    safe_name = inv.doc_number.replace('/', '_')
    return StreamingResponse(
        pdf_buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f"inline; filename=Invoice_{safe_name}.pdf"}
    )

@app.get("/api/v1/compliance/gstr1/json")
def export_gstr1_json(period: str = "092026", tenant: Tenant = Depends(verify_tenant), db: Session = Depends(get_db)):
    data = generate_gstr1_json_payload(db, tenant, period)
    return JSONResponse(content=data, headers={"Content-Disposition": f"attachment; filename=GSTR1_{period}.json"})

@app.get("/api/v1/compliance/gstr1/excel")
def export_gstr1_excel(period: str = "092026", tenant: Tenant = Depends(verify_tenant), db: Session = Depends(get_db)):
    stream = generate_gstr1_excel_workbook(db, tenant, period)
    return Response(
        content=stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=GSTR1_{period}.xlsx"}
    )

# -------------------------------------------------------------
# GSTR-2B DRAG & DROP MULTIPART FILE RECONCILIATION
# -------------------------------------------------------------
class ReconciliationPayload(BaseModel):
    purchase_register: List[Dict[str, Any]]
    gstr2b_records: List[Dict[str, Any]]

@app.post("/api/v1/compliance/reconcile-files")
async def reconcile_uploaded_files(
    purchase_file: UploadFile = File(..., description="Purchase register Excel/CSV"),
    portal_file: UploadFile = File(..., description="GST Portal GSTR-2B JSON"),
    tenant: Tenant = Depends(verify_tenant),
    db: Session = Depends(get_db)
):
    try:
        pr_bytes = await purchase_file.read()
        gstr2b_bytes = await portal_file.read()

        purchase_records = parse_purchase_register_file(pr_bytes, purchase_file.filename)
        portal_records = parse_gstr2b_json_file(gstr2b_bytes)

        if not purchase_records:
            raise HTTPException(status_code=400, detail="Could not extract any valid invoice rows from the purchase file. Please ensure columns include Invoice No and Taxable Value.")
        if not portal_records:
            raise HTTPException(status_code=400, detail="Could not extract valid B2B data from the GSTR-2B JSON file. Ensure it is an authentic portal export.")

        results = run_gstr2b_reconciliation(purchase_records, portal_records)

        db.add(ComplianceAuditLog(
            tenant_id=tenant.id,
            event_type="LIVE_FILE_RECONCILIATION",
            doc_number=f"PR:{len(purchase_records)}_2B:{len(portal_records)}",
            details=f"Eligible: Rs. {results['summary']['total_eligible_itc']:,.2f} | At-Risk: Rs. {results['summary']['total_at_risk_itc']:,.2f}"
        ))
        db.commit()

        # Attach original parsed records so front-end can export Excel easily
        results["_raw_purchase_register"] = purchase_records
        results["_raw_gstr2b_records"] = portal_records
        return results
    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/api/v1/compliance/reconcile-2b")
def reconcile_itc(
    payload: ReconciliationPayload,
    tenant: Tenant = Depends(verify_tenant),
    db: Session = Depends(get_db)
):
    results = run_gstr2b_reconciliation(payload.purchase_register, payload.gstr2b_records)
    db.add(ComplianceAuditLog(
        tenant_id=tenant.id,
        event_type="GSTR2B_RECONCILIATION_PERFORMED",
        doc_number="ITC-AUDIT",
        details=f"Matched: {results['summary']['matched_count']} | Missing: {results['summary']['missing_in_2b_count']} | At-Risk: Rs. {results['summary']['total_at_risk_itc']:,.2f}"
    ))
    db.commit()
    return results

@app.post("/api/v1/compliance/reconcile-2b/excel")
def export_reconcile_excel(
    payload: ReconciliationPayload,
    tenant: Tenant = Depends(verify_tenant)
):
    results = run_gstr2b_reconciliation(payload.purchase_register, payload.gstr2b_records)
    
    wb = Workbook()
    h_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    h_fill_blue = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    h_fill_rose = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")
    h_fill_green = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
    
    # Sheet 1: Executive Summary
    ws_sum = wb.active
    ws_sum.title = "Audit Summary"
    ws_sum.append(["ApexTax Autonomous Statutory GSTR-2B ITC Audit Report", "", ""])
    ws_sum.append(["Generated On", datetime.now().strftime("%d-%m-%Y %H:%M:%S"), ""])
    ws_sum.append([])
    ws_sum.append(["Audit Metric", "Count", "Tax Amount (INR)"])
    for cell in ws_sum[4]:
        cell.font = h_font
        cell.fill = h_fill_blue
    
    ws_sum.append(["Total Purchase Invoices Reviewed", results["summary"]["total_purchases_reviewed"], ""])
    ws_sum.append(["Matched & Eligible ITC (Sec 16(2)(aa) Safe Claim)", results["summary"]["matched_count"], results["summary"]["total_eligible_itc"]])
    ws_sum.append(["Missing in 2B (Defaulting Vendors - HOLD PAYMENT)", results["summary"]["missing_in_2b_count"], results["summary"]["total_at_risk_itc"]])
    ws_sum.append(["Tax / Rate Discrepancies", results["summary"]["mismatch_count"], ""])
    
    # Sheet 2: Missing in 2B (Payment Hold Ledger)
    ws_missing = wb.create_sheet(title="Missing in 2B - Hold Pay")
    ws_missing.append(["Vendor GSTIN", "Vendor Name", "Invoice No", "Date", "Taxable Value", "ITC at Risk", "Mandated Statutory Action"])
    for cell in ws_missing[1]:
        cell.font = h_font
        cell.fill = h_fill_rose
    for item in results["missing_in_2b"]:
        ws_missing.append([item["vendor_gstin"], item["vendor_name"], item["doc_number"], item.get("doc_date", ""), item["taxable_value"], item["itc_at_risk"], item["action"]])

    # Sheet 3: Matched & Eligible
    ws_matched = wb.create_sheet(title="Matched - Safe ITC")
    ws_matched.append(["Vendor GSTIN", "Vendor Name", "Invoice No", "Date", "Taxable Value", "Eligible ITC", "Statutory Status"])
    for cell in ws_matched[1]:
        cell.font = h_font
        cell.fill = h_fill_green
    for item in results["matched"]:
        ws_matched.append([item["vendor_gstin"], item["vendor_name"], item["doc_number"], item.get("doc_date", ""), item["taxable_value"], item["itc_claimed"], item["remarks"]])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    
    return Response(
        content=stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=GSTR2B_Reconciliation_Audit.xlsx"}
    )

@app.get("/api/v1/invoices/audit-trail")
def get_audit(tenant: Tenant = Depends(verify_tenant), db: Session = Depends(get_db)):
    logs = db.query(ComplianceAuditLog).filter(
        ComplianceAuditLog.tenant_id == tenant.id
    ).order_by(ComplianceAuditLog.timestamp.desc()).limit(15).all()
    return {"tenant": tenant.org_name, "audit_events": logs}
# =====================================================================
# MODULE: DRC-01C (RULE 88D) AUTO-RESOLVER & STATUTORY DEFENSE
# =====================================================================

@app.post("/api/v1/compliance/drc01c/resolve")
async def resolve_drc01c(
    ref_no: str = Form(...),
    period: str = Form(...),
    itc_2b: float = Form(...),
    itc_3b: float = Form(...),
    purchase_file: UploadFile = File(None)
):
    excess_claimed = round(itc_3b - itc_2b, 2)
    pct_variance = round((excess_claimed / itc_2b * 100), 2) if itc_2b > 0 else 100.0

    flagged_suppliers_count = 0
    if purchase_file:
        try:
            import pandas as pd
            content = await purchase_file.read()
            df = pd.read_excel(io.BytesIO(content)) if purchase_file.filename.endswith(('.xlsx', '.xls')) else pd.read_csv(io.BytesIO(content))
            flagged_suppliers_count = len(df)
        except Exception:
            flagged_suppliers_count = 0

    part_b_legal_draft = (
        f"In response to Form GST DRC-01C (Reference: {ref_no}) for the tax period {period}, "
        f"the differential input tax credit of ₹{excess_claimed:,.2f} availed in GSTR-3B over GSTR-2B "
        f"represents genuine inward supplies received across staggered tax periods. "
        f"The credit has been availed strictly in compliance with Section 16(2) and within the statutory timeline "
        f"mandated under Section 16(4) of the CGST Act. No blocked or ineligible credit under Section 17(5) "
        f"has been claimed. Detailed invoice reconciliations and supplier verification declarations are retained on record."
    )

    vendor_whatsapp_directive = (
        f"🚨 *OFFICIAL GST NOTICE RECOVERY INTIMATION*\n"
        f"Notice Ref: {ref_no} | Period: {period}\n\n"
        f"Dear Vendor Partner,\n"
        f"The GST portal has issued Form DRC-01C under Rule 88D citing an ITC variance of ₹{excess_claimed:,.2f}. "
        f"Inward records show invoices pending or omitted from your GSTR-1 filing.\n\n"
        f"Please file your GSTR-1 immediately to ensure credit reflection in GSTR-2B. "
        f"Per statutory guidelines, pending invoice payments remain on temporary hold to prevent Section 50 interest liabilities."
    )

    return JSONResponse(content={
        "status": "SUCCESS",
        "notice_summary": {
            "ref_number": ref_no,
            "period": period,
            "gstr2b_itc": itc_2b,
            "gstr3b_itc": itc_3b,
            "excess_difference": excess_claimed,
            "variance_percentage": f"{pct_variance}%",
            "suppliers_analyzed": flagged_suppliers_count
        },
        "part_b_reply": part_b_legal_draft,
        "vendor_whatsapp_directive": vendor_whatsapp_directive
    })
    # =====================================================================
# MODULE: DRC-01B (RULE 88C) OUTWARD LIABILITY AUTO-DEFENSE
# =====================================================================

@app.post("/api/v1/compliance/drc01b/resolve")
async def resolve_drc01b(
    ref_no: str = Form(...),
    period: str = Form(...),
    tax_1: float = Form(...),
    tax_3b: float = Form(...),
    reason_code: str = Form("TYPO_SUBSEQUENT_ADJUSTMENT")
):
    excess_liability = round(tax_1 - tax_3b, 2)
    pct_variance = round((excess_liability / tax_3b * 100), 2) if tax_3b > 0 else 100.0

    reason_templates = {
        "TYPO_SUBSEQUENT_ADJUSTMENT": (
            f"The differential outward liability of ₹{excess_liability:,.2f} reported in Form GSTR-1 over Form GSTR-3B "
            f"arose due to inadvertent typographical data entry in Table 4/Table 5 of GSTR-1. The actual taxable supplies and tax "
            f"collected are accurately discharged in GSTR-3B. The rectification/credit note adjustment is being effected in subsequent "
            f"period returns strictly in terms of Circular No. 170/02/2022-GST."
        ),
        "CANCELLED_EINVOICE": (
            f"The variance of ₹{excess_liability:,.2f} represents e-invoices auto-populated from the IRP into GSTR-1 that were "
            f"subsequently cancelled by mutual consent with the recipient prior to dispatch of goods. The actual supply was not executed, "
            f"and hence no tax was collectible or payable under Section 9 of the CGST Act."
        ),
        "ADVANCE_ADJUSTMENT": (
            f"The differential amount of ₹{excess_liability:,.2f} pertains to advances adjusted against tax invoices issued in the current period, "
            f"where tax had already been paid in preceding tax periods under Section 12/13 of the CGST Act. No revenue loss has occurred."
        )
    }

    legal_reply = reason_templates.get(reason_code, reason_templates["TYPO_SUBSEQUENT_ADJUSTMENT"])
    portal_part_b = (
        f"In response to Form GST DRC-01B (Ref: {ref_no}) for the tax period {period}, "
        f"{legal_reply} Certified reconciliation certificates and transaction audit trails are maintained for verification."
    )

    return JSONResponse(content={
        "status": "SUCCESS",
        "notice_summary": {
            "ref_number": ref_no,
            "period": period,
            "gstr1_tax": tax_1,
            "gstr3b_tax": tax_3b,
            "excess_outward_flagged": excess_liability,
            "variance_percentage": f"{pct_variance}%",
            "statutory_risk": "GSTR-1 Portal Lockout Risk under Rule 59(6)" if excess_liability > 0 else "Compliant"
        },
        "part_b_legal_defense": portal_part_b
    })